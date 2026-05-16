#!/usr/bin/env python
# -*- coding: utf-8 -*-

import re
from docx import Document
from docx.shared import Inches, Pt
import os

def parse_table(rows):
    if len(rows) < 3:
        return None
    headers = [h.strip() for h in rows[0].split('|')[1:-1]]
    table_data = [headers]
    for data_row in rows[2:]:
        cells = [c.strip() for c in data_row.split('|')[1:-1]]
        table_data.append(cells)
    return table_data

def md_to_docx(md_file, docx_file):
    with open(md_file, 'r', encoding='utf-8') as f:
        md_content = f.read()

    doc = Document()
    lines = md_content.split('\n')

    in_code_block = False
    code_content = []
    table_rows = []
    i = 0

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```') and not in_code_block:
            in_code_block = True
            code_content = []
            i += 1
            continue
        elif line.strip().startswith('```') and in_code_block:
            in_code_block = False
            code_para = doc.add_paragraph()
            code_run = code_para.add_run('\n'.join(code_content))
            code_run.font.name = 'Consolas'
            code_run.font.size = Pt(9)
            code_para.paragraph_format.left_indent = Inches(0.3)
            i += 1
            continue

        if in_code_block:
            code_content.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_rows.append(line.strip())
            i += 1
            continue
        elif table_rows:
            table_data = parse_table(table_rows)
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for row_idx, row_data in enumerate(table_data):
                    row_cells = table.rows[row_idx].cells
                    for col_idx, cell_text in enumerate(row_data):
                        cell = row_cells[col_idx]
                        cell.text = cell_text
                        if row_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.font.bold = True
                doc.add_paragraph()
            table_rows = []
            continue

        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.add_run('─' * 60)
        elif line.strip() == '':
            doc.add_paragraph()
        else:
            line_processed = re.sub(r'\*\*([^*]+)\*\*', r'\1', line)
            line_processed = re.sub(r'`([^`]+)`', r'\1', line_processed)
            doc.add_paragraph(line_processed)

        i += 1

    doc.save(docx_file)
    print(f"文档已成功保存到: {docx_file}")

if __name__ == '__main__':
    md_file = r'c:\Users\user\Desktop\毕业设计\毕业论文-夏永基-深度改写版.md'
    docx_file = r'c:\Users\user\Desktop\毕业设计\毕业论文-夏永基-深度改写版.docx'

    if os.path.exists(md_file):
        md_to_docx(md_file, docx_file)
    else:
        print(f"找不到文件: {md_file}")
