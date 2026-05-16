from docx import Document

doc = Document(r'c:\Users\user\Desktop\毕业设计\开题报告-夏永基.docx')

full_text = []

full_text.append("=== 段落内容 ===")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        full_text.append(f"{i+1}: {text}")

full_text.append("\n=== 表格内容 ===")
for i, table in enumerate(doc.tables):
    full_text.append(f"表格 {i+1}:")
    for row in table.rows:
        row_text = []
        for cell in row.cells:
            row_text.append(cell.text.strip())
        full_text.append(" | ".join(row_text))

full_text.append("\n=== 样式信息 ===")
styles = []
for style in doc.styles:
    if style.type == 1:  # 段落样式
        styles.append(style.name)
full_text.append("段落样式: " + ", ".join(styles))

with open(r'c:\Users\user\Desktop\毕业设计\开题报告内容.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(full_text))

print(f"共读取 {len(doc.paragraphs)} 个段落，{len(doc.tables)} 个表格")
print("开题报告内容已保存到文件")