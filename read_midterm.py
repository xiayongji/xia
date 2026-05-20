# -*- coding: utf-8 -*-
from docx import Document
import sys

def read_docx(filepath):
    doc = Document(filepath)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)

    # 读取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text.append(' | '.join(row_text))

    return '\n'.join(text)

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法: python read_docx.py <文件路径>")
        sys.exit(1)

    filepath = sys.argv[1]
    content = read_docx(filepath)
    print(content)
